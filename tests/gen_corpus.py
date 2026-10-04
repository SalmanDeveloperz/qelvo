"""Generate the import test corpus: realistic resumes in many layouts, each with
exact ground truth (the Resume model as JSON), plus negative cases.

    python tests/gen_corpus.py          # writes tests/corpus/*.pdf|docx + *.json

Layouts imitate what people actually send: Google Docs exports, Word documents,
Jake's-style LaTeX, sidebar templates, paragraph descriptions without bullets,
dates in a left column, weak all-caps headings, labelled contact lines, grids.
"""
import json, os, re
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import simpleSplit

OUT = os.path.join(os.path.dirname(__file__), "corpus")
os.makedirs(OUT, exist_ok=True)
# Real-world resumes are mostly set in Microsoft fonts, so the corpus is too. They're
# not redistributable, which is why the corpus is generated locally and git-ignored.
# On macOS or Linux, point QELVO_FONTS at a folder holding arial.ttf, calibri.ttf, etc.
FONTS = os.path.join(os.environ.get("QELVO_FONTS", "C:/Windows/Fonts"), "")
for name, file in [("Arial", "arial"), ("Arial-B", "arialbd"), ("Arial-I", "ariali"), ("Arial-BI", "arialbi"),
                   ("Calibri", "calibri"), ("Calibri-B", "calibrib"), ("Calibri-I", "calibrii"),
                   ("Georgia", "georgia"), ("Georgia-B", "georgiab"), ("Georgia-I", "georgiai"),
                   ("Times", "times"), ("Times-B", "timesbd"), ("Times-I", "timesi"),
                   ("Segoe", "segoeui"), ("Segoe-B", "segoeuib"), ("Segoe-I", "segoeuii")]:
    pdfmetrics.registerFont(TTFont(name, FONTS + file + ".ttf"))

# ── Personas ───────────────────────────────────────────────────────────
ENG = dict(
    name="Jane Doe", headline="Senior Backend Engineer", location="Austin, TX", phone="+1 (512) 555-0147",
    email="jane.doe@example.com", linkedin=("linkedin.com/in/janedoe", "https://linkedin.com/in/janedoe"),
    github=("github.com/janedoe", "https://github.com/janedoe"),
    summary="Backend engineer with 8 years of experience building payment and data platforms in Go and Python. Led the migration of a monolith to event-driven services handling 40k requests per second.",
    experience=[
        dict(company="Stripe", role="Senior Software Engineer", start="Mar 2021", end="Present", location="Remote", bullets=[
            "Designed the idempotency layer for the payouts API, eliminating duplicate transfers across 14 retry paths.",
            "Cut p99 latency of the ledger service from 480 ms to 95 ms by replacing synchronous fan-out with a Kafka-backed outbox and batching writes to PostgreSQL.",
            "Mentored four engineers through their first on-call rotations.",
        ]),
        dict(company="Rackspace Technology", role="Software Engineer", start="Jun 2017", end="Feb 2021", location="San Antonio, TX", bullets=[
            "Built a Terraform module library adopted by 30 internal teams.",
            "Automated certificate rotation for 2,000 customer load balancers.",
            "Wrote the incident review template still used company-wide.",
        ]),
        dict(company="Education First", role="Engineering Intern", start="May 2016", end="Aug 2016", location="Boston, MA", bullets=[
            "Shipped a booking-calendar widget in React used by 12 regional offices.",
            "Experience with on-call: shadowed two incident rotations.",
        ]),
    ],
    education=[dict(school="University of Texas at Austin", degree="B.S. in Computer Science", gpa="GPA 3.8/4.0", start="Aug 2012", end="May 2016")],
    skills=[("Languages", "Go, Python, SQL, TypeScript"), ("Infrastructure", "Kubernetes, Terraform, AWS, PostgreSQL"), ("Practices", "Event sourcing, Observability, Code review")],
    projects=[
        dict(name="ledgerctl", tech="Go, PostgreSQL", link="https://github.com/janedoe/ledgerctl", bullets=["Open-source double-entry ledger CLI with 1.2k GitHub stars."]),
        dict(name="Skills Matrix", tech="React, TypeScript", link="https://github.com/janedoe/skills-matrix", bullets=["Team skills tracker used during quarterly planning at two companies."]),
    ],
    certs=["AWS Certified Solutions Architect – Associate", "Certified Kubernetes Administrator (CKA)"],
)

DATA = dict(
    name="Aisha Khan", headline="Data Analyst", location="Manchester, UK", phone="+44 7700 900123",
    email="aisha.khan@example.co.uk", linkedin=("linkedin.com/in/aishakhan", "https://linkedin.com/in/aishakhan"), github=None,
    summary="Data analyst who turns messy operational data into decisions. Four years of SQL, Python and dashboard work in fintech and consulting.",
    experience=[
        dict(company="Monzo Bank", role="Data Analyst", start="01/2022", end="Present", location="Manchester, UK", bullets=[
            "Built the fraud-loss dashboard reviewed weekly by the risk committee.",
            "Reduced monthly reporting time from three days to four hours with dbt models.",
        ]),
        dict(company="Deloitte", role="Junior Analyst", start="09/2019", end="12/2021", location="London, UK", bullets=[
            "Cleaned and reconciled supplier data for a 40-site retail audit.",
            "Presented findings to client finance teams in fortnightly reviews.",
        ]),
    ],
    education=[
        dict(school="University of Manchester", degree="MSc Data Science", gpa="Distinction", start="2018", end="2019"),
        dict(school="University of Leeds", degree="BSc Mathematics", gpa="First Class Honours", start="2015", end="2018"),
    ],
    skills=[("Analysis", "SQL, Python, pandas, dbt"), ("Visualisation", "Tableau, Looker, Excel"), ("Statistics", "Regression, A/B testing, Forecasting")],
    projects=[],
    certs=["Google Data Analytics Certificate", "Tableau Desktop Specialist"],
)

STUDENT = dict(
    name="Ali Raza", headline="", location="Islamabad, Pakistan", phone="+92 300 1234567",
    email="ali.raza@example.pk", linkedin=None, github=("github.com/aliraza", "https://github.com/aliraza"),
    summary="Final-year software engineering student looking for a backend or platform role. Comfortable with Python, Django and Docker; enjoys debugging production issues.",
    experience=[
        dict(company="Systems Limited", role="Software Engineering Intern", start="Jun 2023", end="Aug 2023", location="Lahore, Pakistan", bullets=[
            "Added pagination and caching to an internal REST API, cutting response time by 60%.",
            "Wrote integration tests for the payroll export service.",
        ]),
    ],
    education=[dict(school="National University of Sciences and Technology", degree="BS Software Engineering", gpa="CGPA 3.6/4.0", start="2020", end="2024")],
    skills=["Python", "Django", "Docker", "PostgreSQL", "Git", "Linux", "REST APIs", "React"],
    projects=[
        dict(name="CampusRide", tech="Django, PostgreSQL, Docker", link="https://github.com/aliraza/campusride", bullets=["Carpooling app for 800 students with route matching and ratings."]),
        dict(name="Exam Scheduler", tech="Python, OR-Tools", link="", bullets=["Constraint solver that produced a clash-free exam timetable for 3,000 students."]),
    ],
    certs=[],
    coursework=["Data Structures", "Operating Systems", "Computer Networks", "Database Systems"],
    awards=[dict(title="Dean's Honor List", subtitle="Top 5% of the cohort", date="2022"), dict(title="NUST Hackathon", subtitle="Second place out of 60 teams", date="Mar 2023")],
)

NURSE = dict(
    name="María José Núñez", headline="Registered Nurse", location="Toronto, Canada", phone="+1 416 555 0199",
    email="maria.nunez@example.ca", linkedin=None, github=None,
    summary="Registered nurse with nine years in acute and post-operative care. Calm under pressure, precise with documentation, and trusted to precept new graduates.",
    experience=[
        dict(company="Toronto General Hospital", role="Registered Nurse", start="2019", end="Current", location="Toronto, Canada", bullets=[
            "Coordinate care for up to six post-operative cardiac patients per shift.",
            "Precept new graduate nurses through their first eight weeks on the unit.",
        ]),
        dict(company="Sunnybrook Health Sciences Centre", role="Nursing Assistant", start="Jan. 2016", end="Dec. 2018", location="Toronto, Canada", bullets=[
            "Supported patient mobility and daily care on a 32-bed surgical ward.",
        ]),
    ],
    education=[dict(school="George Brown College", degree="Diploma in Practical Nursing", gpa="", start="2013", end="2015")],
    skills=["Patient assessment", "Medication administration", "Wound care", "Electronic health records", "Infection control"],
    projects=[],
    certs=["Basic Life Support (BLS)", "Advanced Cardiovascular Life Support (ACLS)"],
)

# ── Truth helpers (mirror the parser's normalisation) ─────────────────
def ndate(start, end):
    s = start.strip()
    e = end.strip()
    if re.match(r"(?i)^(present|current|now)$", e): e = "Present"
    return f"{s} – {e}" if e else s

def title_case(t):
    if t != t.upper(): return t
    small = {"and", "of", "the", "in", "for", "&", "to", "a"}
    return " ".join(w if i and w in small else w[:1].upper() + w[1:] for i, w in enumerate(t.lower().split(" ")))

ROLE = {"summary": "summary", "objective": "summary", "profile": "summary", "professional summary": "summary", "experience": "experience", "work experience": "experience",
        "professional experience": "experience", "employment history": "experience", "education": "education", "skills": "skills",
        "technical skills": "skills", "projects": "projects", "certifications": "certifications", "coursework": "coursework",
        "relevant coursework": "coursework", "awards": "awards", "honors & awards": "awards"}

def sec(title, type_, **kw):
    t = title_case(title)
    d = dict(title=t, role=ROLE[t.lower()], type=type_, column="auto", text="", entries=[], skills=[], items=[])
    d.update(kw)
    return d

def ent(title="", subtitle="", date="", location="", meta="", link="", bullets=()):
    return dict(title=title, subtitle=subtitle, date=date, location=location, meta=meta, link=link, bullets=[dict(text=b, link="") for b in bullets])

def contacts_truth(p, order, labels=False):
    out = []
    for k in order:
        v = p.get(k)
        if not v: continue
        if k == "email": out.append(dict(kind="email", text=v, url=f"mailto:{v}"))
        elif k in ("linkedin", "github"): out.append(dict(kind=k, text=v[0], url=v[1]))
        else: out.append(dict(kind=k, text=v, url=""))
    return out

def resume(p, contacts, sections, headline=True):
    return dict(version=1, template="modern", pages=1, paper="letter", name=p["name"], headline=p["headline"] if headline else "",
                contacts=contacts, sections=sections)

# ── A tiny layout kit on top of reportlab ─────────────────────────────
class Doc:
    def __init__(self, path, size=LETTER, margin=54):
        self.c = canvas.Canvas(path, pagesize=size)
        self.W, self.H = size
        self.m = margin
        self.y = self.H - margin
        self.c.setTitle(os.path.basename(path))

    def need(self, h):
        if self.y - h < self.m:
            self.c.showPage()
            self.y = self.H - self.m

    def text(self, x, s, font, size, align="left", y=None, link=None):
        y = self.y if y is None else y
        self.c.setFont(font, size)
        w = pdfmetrics.stringWidth(s, font, size)
        if align == "center": x = x - w / 2
        if align == "right": x = x - w
        self.c.drawString(x, y, s)
        if link: self.c.linkURL(link, (x, y - 2, x + w, y + size), relative=0)
        return x + w

    def line(self, s, font, size, lead=None, x=None, align="left", link=None):
        lead = lead or size * 1.3
        self.need(lead)
        self.y -= lead
        xx = self.m if x is None else x
        if align == "center": xx = self.W / 2
        if align == "right": xx = self.W - self.m
        self.text(xx, s, font, size, align, link=link)

    def lr(self, left, right, font, size, rfont=None, lead=None, x=None):
        lead = lead or size * 1.35
        self.need(lead)
        self.y -= lead
        self.text(self.m if x is None else x, left, font, size)
        self.text(self.W - self.m, right, rfont or font, size, "right")

    def para(self, s, font, size, x=None, width=None, lead=None, align="left"):
        x = self.m if x is None else x
        width = width or (self.W - self.m - x)
        for ln in simpleSplit(s, font, size, width):
            self.line(ln, font, size, lead, x=x if align == "left" else None, align=align)

    def bullet(self, s, font, size, glyph="•", indent=18, gap=12, gfont=None, lead=None, x=None, width=None):
        x = (self.m if x is None else x) + indent
        width = width or (self.W - self.m - x - gap)
        lines = simpleSplit(s, font, size, width)
        for i, ln in enumerate(lines):
            lead2 = lead or size * 1.35
            self.need(lead2)
            self.y -= lead2
            if i == 0: self.text(x, glyph, gfont or font, size)
            self.text(x + gap, ln, font, size)

    def rule(self, w=0.8, gray=0.0, pad=6):
        self.y -= pad
        self.c.setLineWidth(w)
        self.c.setStrokeGray(gray)
        self.c.line(self.m, self.y, self.W - self.m, self.y)
        self.y -= pad

    def gap(self, h):
        self.y -= h

    def save(self):
        self.c.save()

def write(name, doc_or_none, truth):
    with open(os.path.join(OUT, name + ".json"), "w", encoding="utf-8") as f:
        json.dump(truth, f, indent=1, ensure_ascii=False)

# ── Styles ────────────────────────────────────────────────────────────
def style_docs(p, name, size=A4):
    """Google Docs export: centred name + headline, letter-spaced city, labelled contacts, untitled summary,
    centred caps headings between rules, Role bold + date right, Company bold-italic, ● bullets, skills grid."""
    d = Doc(os.path.join(OUT, name + ".pdf"), size, margin=56)
    d.line(p["name"], "Arial", 30, lead=34, align="center")
    if p["headline"]: d.line(p["headline"], "Arial-B", 17, lead=26, align="center")
    d.line(" ".join(p["location"]), "Arial", 10, lead=20, align="center")
    parts = [("Phone: ", p["phone"], None), ("Email: ", p["email"], "mailto:" + p["email"])]
    if p["linkedin"]: parts.append(("LinkedIn: ", p["linkedin"][0], p["linkedin"][1]))
    d.gap(15)
    widths = sum(pdfmetrics.stringWidth(a, "Arial-B", 10) + pdfmetrics.stringWidth(b, "Arial", 10) for a, b, _ in parts) + 2 * pdfmetrics.stringWidth("  |  ", "Arial", 10)
    x = d.W / 2 - widths / 2
    for i, (lab, val, url) in enumerate(parts):
        if i: x = d.text(x, "  |  ", "Arial", 10)
        x = d.text(x, lab, "Arial-B", 10)
        x = d.text(x, val, "Arial", 10, link=url)
    d.rule(1.5, 0.5, 10)
    d.para(p["summary"], "Arial", 10, lead=14)
    def heading(t):
        d.gap(8); d.rule(1.2, 0, 8); d.line(t, "Arial-B", 12, lead=18, align="center"); d.gap(8)
    heading("WORK EXPERIENCE")
    for e in p["experience"]:
        d.lr(e["role"], f"{e['start']} — {e['end']}", "Arial-B", 12, rfont="Arial-B", lead=20)
        d.line(e["company"], "Arial-BI", 11, lead=15)
        for b in e["bullets"]: d.bullet(b, "Arial", 11, glyph="●", indent=18, gap=22, lead=15)
        d.gap(10)
    heading("SKILLS")
    items = [s for s in p["skills"]] if isinstance(p["skills"][0], str) else [f"{a}: {b}" for a, b in p["skills"]]
    if isinstance(p["skills"][0], str):
        half = (len(items) + 1) // 2
        y0 = d.y
        for col, chunk in enumerate([items[:half], items[half:]]):
            d.y = y0
            for it in chunk: d.bullet(it, "Arial", 11, glyph="●", indent=20 + col * 240, gap=22, lead=19, width=200)
        d.y = y0 - 19 * half
    else:
        for it in items: d.bullet(it, "Arial", 11, glyph="●", indent=20, gap=22, lead=17)
    heading("QUALIFICATION")
    for ed in p["education"]:
        deg = ed["degree"] + (f" ({ed['gpa']})" if ed["gpa"] and "/" in ed["gpa"] else "")
        d.lr(deg, f"{ed['start']} – {ed['end']}", "Arial", 11, rfont="Arial-B", lead=18)
        d.line(ed["school"], "Arial-B", 11, lead=15)
        d.gap(6)
    d.save()
    skills_sec = sec("SKILLS", "list", items=[dict(text=i, link="") for i in items]) if isinstance(p["skills"][0], str) else \
        sec("SKILLS", "skills", skills=[dict(label=a, value=b) for a, b in p["skills"]])
    edu = [ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"]), meta=ed["gpa"] if "/" in ed["gpa"] else "") for ed in p["education"]]
    truth = resume(p, contacts_truth(p, ["location"]) + contacts_truth(p, ["phone", "email", "linkedin"]), [
        sec("SUMMARY", "summary", text=p["summary"]),
        sec("WORK EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), bullets=e["bullets"]) for e in p["experience"]]),
        skills_sec,
        dict(sec("EDUCATION", "entries"), title="Qualification", entries=edu),
    ])
    write(name, d, truth)

def style_word(p, name):
    """Word: left-aligned, bold caps headings with an underline, "Company, City" bold + dates right,
    role italic, en-dash bullets, labelled skills lines, plain contact line with | separators."""
    d = Doc(os.path.join(OUT, name + ".pdf"), LETTER, margin=54)
    d.line(p["name"].upper(), "Calibri-B", 22, lead=24)
    bits = [p["location"], p["phone"], p["email"]] + ([p["linkedin"][0]] if p["linkedin"] else []) + ([p["github"][0]] if p["github"] else [])
    d.gap(16)
    x = d.m
    links = {p["email"]: "mailto:" + p["email"]}
    if p["linkedin"]: links[p["linkedin"][0]] = p["linkedin"][1]
    if p["github"]: links[p["github"][0]] = p["github"][1]
    for i, b in enumerate(bits):
        if i: x = d.text(x, " | ", "Calibri", 10.5)
        x = d.text(x, b, "Calibri", 10.5, link=links.get(b))
    def heading(t):
        d.gap(10); d.line(t, "Calibri-B", 12.5, lead=14); d.rule(0.6, 0, 3)
    heading("PROFESSIONAL SUMMARY")
    d.para(p["summary"], "Calibri", 10.5, lead=13.5)
    heading("EXPERIENCE")
    for e in p["experience"]:
        d.lr(f"{e['company']}, {e['location']}", f"{e['start']} – {e['end']}", "Calibri-B", 11, lead=17)
        d.line(e["role"], "Calibri-I", 11, lead=14)
        for b in e["bullets"]: d.bullet(b, "Calibri", 10.5, glyph="–", indent=10, gap=10, lead=13.5)
    heading("EDUCATION")
    for ed in p["education"]:
        d.lr(ed["school"], f"{ed['start']} – {ed['end']}", "Calibri-B", 11, lead=17)
        d.line(ed["degree"] + (f", {ed['gpa']}" if ed["gpa"] else ""), "Calibri-I", 11, lead=14)
    heading("TECHNICAL SKILLS")
    for a, b in p["skills"]:
        d.need(14); d.y -= 14
        x = d.text(d.m, a + ": ", "Calibri-B", 10.5)
        d.text(x, b, "Calibri", 10.5)
    if p["projects"]:
        heading("PROJECTS")
        for pr in p["projects"]:
            d.need(17); d.y -= 17
            x = d.text(d.m, pr["name"], "Calibri-B", 11)
            x = d.text(x, " | " + pr["tech"], "Calibri-I", 11)
            if pr["link"]: d.text(x + 6, "repo", "Calibri", 10, link=pr["link"])
            for b in pr["bullets"]: d.bullet(b, "Calibri", 10.5, glyph="–", indent=10, gap=10, lead=13.5)
    if p["certs"]:
        heading("CERTIFICATIONS")
        for c in p["certs"]: d.bullet(c, "Calibri", 10.5, glyph="–", indent=10, gap=10, lead=13.5)
    d.save()
    sections = [sec("PROFESSIONAL SUMMARY", "summary", text=p["summary"])]
    sections[0]["role"] = "summary"
    sections.append(sec("EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), e["location"], bullets=e["bullets"]) for e in p["experience"]]))
    sections.append(sec("EDUCATION", "entries", entries=[ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"]), meta=ed["gpa"]) for ed in p["education"]]))
    sections.append(sec("TECHNICAL SKILLS", "skills", skills=[dict(label=a, value=b) for a, b in p["skills"]]))
    if p["projects"]:
        sections.append(sec("PROJECTS", "entries", entries=[ent(pr["name"], meta=pr["tech"], link=pr["link"], bullets=pr["bullets"]) for pr in p["projects"]]))
    if p["certs"]:
        sections.append(sec("CERTIFICATIONS", "list", items=[dict(text=c, link="") for c in p["certs"]]))
    write(name, d, resume(p, contacts_truth(p, ["location", "phone", "email", "linkedin", "github"]), sections, headline=False))

def style_jake(p, name):
    """Jake's resume (LaTeX look): Company bold + location right, role italic + dates italic right, • bullets."""
    d = Doc(os.path.join(OUT, name + ".pdf"), LETTER, margin=40)
    d.line(p["name"], "Times-B", 24, lead=26, align="center")
    bits = [p["phone"], p["email"]] + ([p["linkedin"][0]] if p["linkedin"] else []) + ([p["github"][0]] if p["github"] else [])
    d.line(" | ".join(bits), "Times", 10, lead=15, align="center")
    def heading(t):
        d.gap(6); d.line(t, "Times-B", 12, lead=14); d.rule(0.4, 0, 2)
    heading("EDUCATION")
    for ed in p["education"]:
        d.lr(ed["school"], p["location"], "Times-B", 11, rfont="Times", lead=15)
        d.lr(ed["degree"] + (f", {ed['gpa']}" if ed["gpa"] else ""), f"{ed['start']} – {ed['end']}", "Times-I", 10, lead=12)
    heading("EXPERIENCE")
    for e in p["experience"]:
        d.lr(e["company"], e["location"], "Times-B", 11, rfont="Times", lead=15)
        d.lr(e["role"], f"{e['start']} – {e['end']}", "Times-I", 10, lead=12)
        for b in e["bullets"]: d.bullet(b, "Times", 10, glyph="•", indent=16, gap=9, lead=12)
    heading("PROJECTS")
    for pr in p["projects"]:
        d.need(15); d.y -= 15
        x = d.text(d.m, pr["name"], "Times-B", 10.5)
        x = d.text(x, " | ", "Times", 10.5)
        x = d.text(x, pr["tech"], "Times-I", 10.5)
        if pr["link"]:
            d.c.linkURL(pr["link"], (x + 4, d.y - 2, x + 14, d.y + 9), relative=0)
        for b in pr["bullets"]: d.bullet(b, "Times", 10, glyph="•", indent=16, gap=9, lead=12)
    if p.get("awards"):
        heading("AWARDS")
        for a in p["awards"]:
            d.lr(a["title"], a["date"], "Times-B", 10.5, rfont="Times-I", lead=14)
            d.line(a["subtitle"], "Times-I", 10, lead=12)
    heading("TECHNICAL SKILLS")
    d.para(", ".join(p["skills"]), "Times", 10, lead=12)
    if p.get("coursework"):
        heading("RELEVANT COURSEWORK")
        for c in p["coursework"]: d.bullet(c, "Times", 10, glyph="•", indent=16, gap=9, lead=12)
    d.save()
    loc = p["location"]
    sections = [
        sec("EDUCATION", "entries", entries=[ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"]), loc, meta=ed["gpa"]) for ed in p["education"]]),
        sec("EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), e["location"], bullets=e["bullets"]) for e in p["experience"]]),
        sec("PROJECTS", "entries", entries=[ent(pr["name"], meta=pr["tech"], link=pr["link"], bullets=pr["bullets"]) for pr in p["projects"]]),
    ]
    if p.get("awards"): sections.append(sec("AWARDS", "entries", entries=[ent(a["title"], a["subtitle"], a["date"]) for a in p["awards"]]))
    sections.append(sec("TECHNICAL SKILLS", "list", items=[dict(text=s, link="") for s in p["skills"]]))
    if p.get("coursework"): sections.append(sec("RELEVANT COURSEWORK", "list", items=[dict(text=c, link="") for c in p["coursework"]]))
    write(name, d, resume(p, contacts_truth(p, ["phone", "email", "linkedin", "github"]), sections, headline=False))

def style_sidebar(p, name):
    """Canva-style: name across the top, a narrow left sidebar (Contact, Skills, Education), main column
    on the right (Profile, Experience). Role | Company on one line, dates under it."""
    d = Doc(os.path.join(OUT, name + ".pdf"), A4, margin=40)
    c = d.c
    d.line(p["name"], "Segoe-B", 26, lead=30)
    if p["headline"]: d.line(p["headline"], "Segoe", 13, lead=18)
    top = d.y - 18
    c.setLineWidth(0.5); c.line(d.m, top + 8, d.W - d.m, top + 8)
    side_x, side_w, main_x = d.m, 150, d.m + 180
    # sidebar
    d.y = top
    def sh(t, x):
        d.gap(6); d.line(t, "Segoe-B", 11, lead=16, x=x)
    sh("CONTACT", side_x)
    for k in ["phone", "email", "location"]:
        d.line(p[k], "Segoe", 9, lead=13, x=side_x, link=("mailto:" + p[k]) if k == "email" else None)
    if p["linkedin"]: d.line(p["linkedin"][0], "Segoe", 9, lead=13, x=side_x, link=p["linkedin"][1])
    sh("SKILLS", side_x)
    items = p["skills"] if isinstance(p["skills"][0], str) else [i.strip() for _, v in p["skills"] for i in v.split(",")]
    for it in items: d.bullet(it, "Segoe", 9, glyph="•", indent=0, gap=9, lead=13, x=side_x, width=side_w - 10)
    sh("EDUCATION", side_x)
    for ed in p["education"]:
        d.para(ed["degree"], "Segoe-B", 9, x=side_x, width=side_w, lead=12)
        d.para(ed["school"], "Segoe", 9, x=side_x, width=side_w, lead=12)
        d.line(f"{ed['start']} – {ed['end']}", "Segoe-I", 9, lead=12, x=side_x)
    # main
    d.y = top
    sh("PROFILE", main_x)
    d.para(p["summary"], "Segoe", 9.5, x=main_x, lead=13)
    sh("EXPERIENCE", main_x)
    for e in p["experience"]:
        d.gap(4)
        d.line(f"{e['role']} | {e['company']}", "Segoe-B", 10, lead=14, x=main_x)
        d.line(f"{e['start']} – {e['end']} | {e['location']}", "Segoe-I", 9, lead=12, x=main_x)
        for b in e["bullets"]: d.bullet(b, "Segoe", 9.5, glyph="•", indent=4, gap=9, lead=13, x=main_x)
    d.save()
    sections = [
        sec("SKILLS", "list", items=[dict(text=i, link="") for i in items]),
        sec("EDUCATION", "entries", entries=[ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"])) for ed in p["education"]]),
        dict(sec("PROFILE", "summary", text=p["summary"])),
        sec("EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), e["location"], bullets=e["bullets"]) for e in p["experience"]]),
    ]
    write(name, d, resume(p, contacts_truth(p, ["phone", "email", "location", "linkedin"]), sections))

def style_paragraph(p, name):
    """No bullet glyphs: each role has a description paragraph (wrapped lines), "Role at Company" bold,
    a meta line with dates · location, plain Georgia headings in title case."""
    d = Doc(os.path.join(OUT, name + ".pdf"), A4, margin=60)
    d.line(p["name"], "Georgia-B", 22, lead=26)
    d.line(f"{p['email']}  ·  {p['phone']}  ·  {p['location']}", "Georgia", 9.5, lead=16)
    def heading(t):
        d.gap(12); d.line(t, "Georgia-B", 13, lead=16); d.gap(2)
    heading("Summary")
    d.para(p["summary"], "Georgia", 10, lead=14)
    heading("Experience")
    for e in p["experience"]:
        d.gap(4)
        d.line(f"{e['role']} at {e['company']}", "Georgia-B", 10.5, lead=15)
        d.line(f"{e['start']} – {e['end']} · {e['location']}", "Georgia-I", 9.5, lead=13)
        d.para(" ".join(e["bullets"]), "Georgia", 10, lead=14)
    heading("Education")
    for ed in p["education"]:
        d.line(ed["school"], "Georgia-B", 10.5, lead=15)
        d.line(f"{ed['degree']} · {ed['start']} – {ed['end']}", "Georgia", 10, lead=13)
    heading("Skills")
    for a, b in p["skills"]: d.line(f"{a}: {b}", "Georgia", 10, lead=14)
    d.save()
    sections = [
        sec("SUMMARY", "summary", text=p["summary"]),
        sec("EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), e["location"], bullets=[" ".join(e["bullets"])]) for e in p["experience"]]),
        sec("EDUCATION", "entries", entries=[ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"])) for ed in p["education"]]),
        sec("SKILLS", "skills", skills=[dict(label=a, value=b) for a, b in p["skills"]]),
    ]
    for s_ in sections: s_["title"] = s_["title"]
    write(name, d, resume(p, contacts_truth(p, ["email", "phone", "location"]), sections, headline=False))

def style_dates_left(p, name):
    """Dates in a left column, role/company to the right on the same baseline; small caps-ish headings."""
    d = Doc(os.path.join(OUT, name + ".pdf"), LETTER, margin=50)
    d.line(p["name"], "Cambria" if False else "Times-B", 22, lead=24)
    if p["headline"]: d.line(p["headline"], "Times-I", 12, lead=16)
    d.line(f"{p['location']} | {p['phone']} | {p['email']}", "Times", 10, lead=15)
    col = d.m + 105
    def heading(t):
        d.gap(10); d.line(t, "Times-B", 12, lead=15); d.rule(0.5, 0.4, 3)
    heading("PROFILE")
    d.para(p["summary"], "Times", 10.5, lead=14)
    heading("EXPERIENCE")
    for e in p["experience"]:
        d.need(16); d.y -= 16
        d.text(d.m, f"{e['start']} – {e['end']}", "Times", 10)
        d.text(col, f"{e['role']}, {e['company']}", "Times-B", 10.5)
        d.line(e["location"], "Times-I", 10, lead=13, x=col)
        for b in e["bullets"]: d.bullet(b, "Times", 10.5, glyph="•", indent=105, gap=10, lead=13.5)
    heading("EDUCATION")
    for ed in p["education"]:
        d.need(16); d.y -= 16
        d.text(d.m, f"{ed['start']} – {ed['end']}", "Times", 10)
        d.text(col, ed["degree"], "Times-B", 10.5)
        d.line(ed["school"], "Times-I", 10, lead=13, x=col)
    heading("SKILLS")
    d.para(", ".join(p["skills"]), "Times", 10.5, lead=14)
    if p["certs"]:
        heading("CERTIFICATIONS")
        for cc in p["certs"]: d.bullet(cc, "Times", 10.5, glyph="•", indent=0, gap=10, lead=13.5)
    d.save()
    sections = [
        sec("PROFILE", "summary", text=p["summary"]),
        sec("EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), e["location"], bullets=e["bullets"]) for e in p["experience"]]),
        sec("EDUCATION", "entries", entries=[ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"])) for ed in p["education"]]),
        sec("SKILLS", "list", items=[dict(text=s, link="") for s in p["skills"]]),
    ]
    if p["certs"]: sections.append(sec("CERTIFICATIONS", "list", items=[dict(text=c, link="") for c in p["certs"]]))
    write(name, d, resume(p, contacts_truth(p, ["location", "phone", "email"]), sections))

def style_weak(p, name):
    """Weak signals: headings are plain caps at body size (not bold), one-line "Company — Role — Dates"
    headers in bold, '*' bullets, Courier-ish plainness via Arial."""
    d = Doc(os.path.join(OUT, name + ".pdf"), A4, margin=50)
    d.line(p["name"], "Arial-B", 16, lead=18)
    bits = [p["email"], p["phone"], p["location"]] + ([p["github"][0]] if p["github"] else [])
    d.line(" • ".join(bits), "Arial", 9.5, lead=14)
    def heading(t):
        d.gap(10); d.line(t, "Arial", 10, lead=14); d.gap(1)
    heading("OBJECTIVE")
    d.para(p["summary"], "Arial", 9.5, lead=13)
    heading("EXPERIENCE")
    for e in p["experience"]:
        d.line(f"{e['company']} — {e['role']} — {e['start']} – {e['end']}", "Arial-B", 9.5, lead=15)
        for b in e["bullets"]: d.bullet(b, "Arial", 9.5, glyph="*", indent=8, gap=8, lead=12.5)
    heading("EDUCATION")
    for ed in p["education"]:
        d.line(f"{ed['school']} — {ed['degree']} — {ed['start']} – {ed['end']}", "Arial-B", 9.5, lead=15)
        if ed["gpa"]: d.line(ed["gpa"], "Arial", 9.5, lead=12.5)
    heading("PROJECTS")
    for pr in p["projects"]:
        d.line(f"{pr['name']} — {pr['tech']}", "Arial-B", 9.5, lead=15)
        for b in pr["bullets"]: d.bullet(b, "Arial", 9.5, glyph="*", indent=8, gap=8, lead=12.5)
    heading("SKILLS")
    d.para(", ".join(p["skills"]), "Arial", 9.5, lead=13)
    d.save()
    sections = [
        sec("OBJECTIVE", "summary", text=p["summary"]),
        sec("EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), bullets=e["bullets"]) for e in p["experience"]]),
        sec("EDUCATION", "entries", entries=[ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"]), meta=ed["gpa"]) for ed in p["education"]]),
        sec("PROJECTS", "entries", entries=[ent(pr["name"], meta=pr["tech"], bullets=pr["bullets"]) for pr in p["projects"]]),
        sec("SKILLS", "list", items=[dict(text=s, link="") for s in p["skills"]]),
    ]
    write(name, d, resume(p, contacts_truth(p, ["email", "phone", "location", "github"]), sections, headline=False))

# ── DOCX ──────────────────────────────────────────────────────────────
def docx_styled(p, name):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.add_heading(p["name"], level=0)
    if p["headline"]: doc.add_paragraph().add_run(p["headline"]).bold = True
    doc.add_paragraph(" | ".join([p["location"], p["phone"], p["email"]]))
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(p["summary"])
    doc.add_heading("Experience", level=1)
    for e in p["experience"]:
        para = doc.add_paragraph()
        para.add_run(f"{e['role']}, {e['company']}").bold = True
        para.add_run(f"\t{e['start']} – {e['end']}")
        doc.add_paragraph().add_run(e["location"]).italic = True
        for b in e["bullets"]: doc.add_paragraph(b, style="List Bullet")
    doc.add_heading("Education", level=1)
    for ed in p["education"]:
        para = doc.add_paragraph()
        para.add_run(ed["school"]).bold = True
        para.add_run(f"\t{ed['start']} – {ed['end']}")
        doc.add_paragraph(ed["degree"] + (f" ({ed['gpa']})" if ed["gpa"] else ""))
    doc.add_heading("Skills", level=1)
    items = p["skills"] if isinstance(p["skills"][0], str) else None
    if items:
        for s in items: doc.add_paragraph(s, style="List Bullet")
    else:
        for a, b in p["skills"]:
            para = doc.add_paragraph(); para.add_run(a + ": ").bold = True; para.add_run(b)
    if p["certs"]:
        doc.add_heading("Certifications", level=1)
        for c in p["certs"]: doc.add_paragraph(c, style="List Bullet")
    doc.save(os.path.join(OUT, name + ".docx"))
    sections = [
        sec("SUMMARY", "summary", text=p["summary"]),
        sec("EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), e["location"], bullets=e["bullets"]) for e in p["experience"]]),
        sec("EDUCATION", "entries", entries=[ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"]), meta=ed["gpa"]) for ed in p["education"]]),
        sec("SKILLS", "list", items=[dict(text=s, link="") for s in items]) if items else sec("SKILLS", "skills", skills=[dict(label=a, value=b) for a, b in p["skills"]]),
    ]
    if p["certs"]: sections.append(sec("CERTIFICATIONS", "list", items=[dict(text=c, link="") for c in p["certs"]]))
    write(name, None, resume(p, contacts_truth(p, ["location", "phone", "email"]), sections))

def docx_plain(p, name):
    """No Word heading styles at all: bold caps paragraphs as headings, '•' typed by hand."""
    from docx import Document
    doc = Document()
    doc.add_paragraph().add_run(p["name"].upper()).bold = True
    doc.add_paragraph(f"{p['email']} • {p['phone']} • {p['location']}")
    def h(t): doc.add_paragraph().add_run(t).bold = True
    h("OBJECTIVE"); doc.add_paragraph(p["summary"])
    h("EDUCATION")
    for ed in p["education"]:
        doc.add_paragraph().add_run(ed["school"]).bold = True
        doc.add_paragraph(f"{ed['degree']}, {ed['gpa']}, {ed['start']} – {ed['end']}")
    h("PROJECTS")
    for pr in p["projects"]:
        para = doc.add_paragraph(); para.add_run(pr["name"]).bold = True; para.add_run(f" | {pr['tech']}")
        for b in pr["bullets"]: doc.add_paragraph("• " + b)
    h("EXPERIENCE")
    for e in p["experience"]:
        doc.add_paragraph().add_run(f"{e['company']} | {e['role']} | {e['start']} – {e['end']}").bold = True
        for b in e["bullets"]: doc.add_paragraph("• " + b)
    h("SKILLS"); doc.add_paragraph(", ".join(p["skills"]))
    doc.save(os.path.join(OUT, name + ".docx"))
    sections = [
        sec("OBJECTIVE", "summary", text=p["summary"]),
        sec("EDUCATION", "entries", entries=[ent(ed["school"], ed["degree"], ndate(ed["start"], ed["end"]), meta=ed["gpa"]) for ed in p["education"]]),
        sec("PROJECTS", "entries", entries=[ent(pr["name"], meta=pr["tech"], bullets=pr["bullets"]) for pr in p["projects"]]),
        sec("EXPERIENCE", "entries", entries=[ent(e["company"], e["role"], ndate(e["start"], e["end"]), bullets=e["bullets"]) for e in p["experience"]]),
        sec("SKILLS", "list", items=[dict(text=s, link="") for s in p["skills"]]),
    ]
    write(name, None, resume(dict(p, name=p["name"]), contacts_truth(p, ["email", "phone", "location"]), sections, headline=False))

# ── Negative cases ────────────────────────────────────────────────────
def neg_blank():
    c = canvas.Canvas(os.path.join(OUT, "neg_blank.neg.pdf"), pagesize=A4); c.showPage(); c.save()

def neg_article():
    d = Doc(os.path.join(OUT, "neg_article.neg.pdf"), A4)
    d.line("Why Distributed Systems Are Hard", "Georgia-B", 20, lead=24)
    for _ in range(3):
        d.para("Networks partition, clocks drift and processes pause at the worst possible moment. A system that tolerates these failures has to be designed for them from the start, not patched after the first outage. " * 2, "Georgia", 11, lead=15)
        d.gap(8)
    d.save()

if __name__ == "__main__":
    style_docs(ENG, "docs_eng")
    style_docs(STUDENT, "docs_student", size=A4)
    style_docs(NURSE, "docs_nurse")
    style_word(DATA, "word_data")
    style_word(ENG, "word_eng")
    style_jake(STUDENT, "jake_student")
    style_sidebar(ENG, "sidebar_eng")
    style_sidebar(DATA, "sidebar_data")
    style_paragraph(DATA, "paragraph_data")
    style_dates_left(NURSE, "datesleft_nurse")
    style_weak(STUDENT, "weak_student")
    docx_styled(DATA, "docx_data")
    docx_styled(NURSE, "docx_nurse")
    docx_plain(STUDENT, "docx_plain_student")
    neg_blank()
    neg_article()
    print("corpus:", sorted(os.listdir(OUT)))
